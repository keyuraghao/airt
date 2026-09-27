"""Render a Report into any supported output format.

render(report, fmt) -> (payload bytes, media type, file extension)
"""
from __future__ import annotations

import csv
import io
import json
import math
import re
import xml.etree.ElementTree as ET
from collections.abc import Callable
from datetime import UTC, datetime
from html import escape
from typing import Any

import yaml

from .. import __version__
from .model import ChartSection, FindingsSection, KeyValueSection, Report, Section, TableSection, TextSection

FORMATS: dict[str, dict[str, str]] = {
    "json": {"media_type": "application/json", "extension": "json", "description": "Report document as JSON"},
    "yaml": {"media_type": "application/yaml", "extension": "yaml", "description": "Report document as YAML"},
    "csv": {"media_type": "text/csv", "extension": "csv", "description": "Comma separated values, one block per section"},
    "tsv": {"media_type": "text/tab-separated-values", "extension": "tsv", "description": "Tab separated values, one block per section"},
    "md": {"media_type": "text/markdown", "extension": "md", "description": "Markdown"},
    "html": {"media_type": "text/html", "extension": "html", "description": "Self-contained printable HTML with inline charts"},
    "pdf": {"media_type": "application/pdf", "extension": "pdf", "description": "PDF document"},
    "xlsx": {"media_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "extension": "xlsx", "description": "Excel workbook, one sheet per section"},
    "txt": {"media_type": "text/plain", "extension": "txt", "description": "Plain text with aligned tables"},
    "xml": {"media_type": "application/xml", "extension": "xml", "description": "Generic XML of the report document"},
    "sarif": {"media_type": "application/sarif+json", "extension": "sarif", "description": "SARIF 2.1.0 static analysis results (findings and vulnerable probes)"},
    "junit": {"media_type": "application/xml", "extension": "xml", "description": "JUnit XML for CI (each probe result or ticket is a testcase)"},
}

ALIASES: dict[str, str] = {"yml": "yaml", "markdown": "md", "text": "txt", "excel": "xlsx", "junitxml": "junit", "sarif.json": "sarif"}
MEDIA_TO_FORMAT: dict[str, str] = {
    "application/json": "json",
    "text/json": "json",
    "application/yaml": "yaml",
    "application/x-yaml": "yaml",
    "text/yaml": "yaml",
    "text/csv": "csv",
    "text/tab-separated-values": "tsv",
    "text/markdown": "md",
    "text/html": "html",
    "application/xhtml+xml": "html",
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "application/vnd.ms-excel": "xlsx",
    "text/plain": "txt",
    "application/xml": "xml",
    "text/xml": "xml",
    "application/sarif+json": "sarif",
    "application/junit+xml": "junit",
}

# Validated categorical palette (light and dark surfaces) plus text tokens.
PALETTE_LIGHT: list[str] = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#4a3aa7", "#e34948"]
PALETTE_DARK: list[str] = ["#3987e5", "#d95926", "#199e70", "#c98500", "#d55181", "#008300", "#9085e9", "#e66767"]
SEVERITY_TO_SARIF: dict[str, str] = {"CRITICAL": "error", "HIGH": "error", "MEDIUM": "warning", "LOW": "note", "INFO": "note", "NONE": "none"}
FAIL_VALUES: set[str] = {"VULNERABLE", "CRITICAL", "HIGH"}
ERROR_VALUES: set[str] = {"ERROR", "FAILED"}


class UnknownFormat(ValueError):
    pass


def normalize_format(fmt: str | None) -> str:
    key = (fmt or "json").strip().lower().lstrip(".")
    key = ALIASES.get(key, key)
    if key not in FORMATS:
        raise UnknownFormat(f"unknown report format '{fmt}'; supported: {', '.join(FORMATS)}")
    return key


def negotiate(accept_header: str | None, explicit_format: str | None = None) -> str:
    """Explicit format wins; otherwise pick the best supported media type from Accept; default json."""
    if explicit_format:
        return normalize_format(explicit_format)
    if not accept_header:
        return "json"
    candidates: list[tuple[float, int, str]] = []
    for index, part in enumerate(accept_header.split(",")):
        pieces = [p.strip() for p in part.split(";") if p.strip()]
        if not pieces:
            continue
        media = pieces[0].lower()
        q = 1.0
        for p in pieces[1:]:
            if p.startswith("q="):
                try:
                    q = float(p[2:])
                except ValueError:
                    q = 0.0
        if q <= 0:
            continue
        fmt = MEDIA_TO_FORMAT.get(media)
        if fmt is None and media in ("*/*", "application/*", "text/*"):
            fmt = "json" if media != "text/*" else "txt"
        if fmt:
            candidates.append((q, -index, fmt))
    if not candidates:
        return "json"
    candidates.sort(reverse=True)
    return candidates[0][2]


# --- helpers ------------------------------------------------------------------------
def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, float):
        return f"{value:.2f}".rstrip("0").rstrip(".") if not value.is_integer() else str(int(value))
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, default=str, ensure_ascii=False)
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    return str(value)


def _plain(value: Any) -> Any:
    """A scalar suitable for JSON/YAML/XLSX cells."""
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    if isinstance(value, datetime):
        return value.isoformat()
    return _cell(value)


def _report_json(report: Report) -> dict[str, Any]:
    return json.loads(json.dumps(report.to_dict(), default=str))


def _slug(text: str, limit: int = 48) -> str:
    s = re.sub(r"[^a-zA-Z0-9]+", "-", text).strip("-").lower()
    return s[:limit] or "report"


def _truncate(text: str, limit: int) -> str:
    return text if len(text) <= limit else text[: max(1, limit - 3)] + "..."


# --- json / yaml --------------------------------------------------------------------
def render_json(report: Report) -> bytes:
    return json.dumps(report.to_dict(), indent=2, default=str, ensure_ascii=False).encode("utf-8")


def render_yaml(report: Report) -> bytes:
    return yaml.safe_dump(_report_json(report), sort_keys=False, allow_unicode=True, default_flow_style=False).encode("utf-8")


# --- csv / tsv ----------------------------------------------------------------------
def _delimited(report: Report, delimiter: str) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=delimiter, lineterminator="\n")
    sections = [s for s in report.sections if s.as_table()[1] or isinstance(s, (TableSection, FindingsSection))]
    single = len(sections) == 1
    if not single:
        writer.writerow(["# " + report.title, report.subtitle, report.generated_at.isoformat(timespec="seconds")])
    for s in sections:
        columns, rows = s.as_table()
        if not single:
            writer.writerow([])
            writer.writerow(["## " + s.title])
        writer.writerow(columns)
        for r in rows:
            writer.writerow([_cell(c) for c in r])
    return buf.getvalue().encode("utf-8")


def render_csv(report: Report) -> bytes:
    return _delimited(report, ",")


def render_tsv(report: Report) -> bytes:
    return _delimited(report, "\t")


# --- text / markdown ----------------------------------------------------------------
def _aligned_table(columns: list[str], rows: list[list[Any]], max_width: int = 40) -> str:
    if not columns:
        return "(empty)"
    cells = [[_truncate(_cell(c).replace("\n", " "), max_width) for c in r] for r in rows]
    widths = [min(max_width, max([len(columns[i])] + [len(r[i]) for r in cells if i < len(r)])) for i in range(len(columns))]
    inner = sum(w + 2 for w in widths) + len(widths) - 1
    if not rows and inner < 11:
        widths[-1] += 11 - inner
    sep = "+" + "+".join("-" * (w + 2) for w in widths) + "+"
    lines = [sep, "|" + "|".join(f" {columns[i].ljust(widths[i])} " for i in range(len(columns))) + "|", sep]
    for r in cells:
        padded = r + [""] * (len(columns) - len(r))
        lines.append("|" + "|".join(f" {padded[i].ljust(widths[i])} " for i in range(len(columns))) + "|")
    lines.append(sep)
    if not rows:
        lines.insert(3, "|" + " (no rows)".ljust(len(sep) - 2) + "|")
    return "\n".join(lines)


def render_txt(report: Report) -> bytes:
    out: list[str] = [report.title.upper(), "=" * len(report.title)]
    if report.subtitle:
        out.append(report.subtitle)
    out.append(f"Generated {report.generated_at.isoformat(timespec='seconds')} by {report.generated_by}")
    for s in report.sections:
        out.append("")
        out.append(s.title)
        out.append("-" * len(s.title))
        if isinstance(s, TextSection):
            out.append(s.text)
        elif isinstance(s, KeyValueSection):
            width = max((len(k) for k, _ in s.rows), default=0)
            out.extend(f"{k.ljust(width)} : {_cell(v)}" for k, v in s.rows)
        else:
            columns, rows = s.as_table()
            if isinstance(s, ChartSection):
                out.append(f"[{s.chart} chart, total {_cell(s.total)}]")
            out.append(_aligned_table(columns, rows))
            if isinstance(s, TableSection) and s.notes:
                out.append(s.notes)
    out.append("")
    return "\n".join(out).encode("utf-8")


def _md_escape(text: str) -> str:
    return text.replace("|", "\\|").replace("\n", " ")


def _md_table(columns: list[str], rows: list[list[Any]]) -> str:
    if not columns:
        return "_(empty)_"
    head = "| " + " | ".join(_md_escape(c) for c in columns) + " |"
    line = "|" + "|".join(" --- " for _ in columns) + "|"
    body = ["| " + " | ".join(_md_escape(_cell(c)) for c in r) + " |" for r in rows]
    if not body:
        body = ["| " + " | ".join("" for _ in columns) + " |"]
    return "\n".join([head, line, *body])


def render_md(report: Report) -> bytes:
    out: list[str] = [f"# {report.title}"]
    if report.subtitle:
        out.append(f"_{report.subtitle}_")
    out.append(f"Generated {report.generated_at.isoformat(timespec='seconds')} by {report.generated_by}")
    for s in report.sections:
        out.append("")
        out.append(f"## {s.title}")
        if isinstance(s, TextSection):
            out.append(s.text if s.markdown else "```\n" + s.text + "\n```")
        else:
            columns, rows = s.as_table()
            out.append(_md_table(columns, rows))
            if isinstance(s, TableSection) and s.notes:
                out.append("")
                out.append(f"_{s.notes}_")
    out.append("")
    return "\n".join(out).encode("utf-8")


# --- html ---------------------------------------------------------------------------
HTML_CSS = """
:root{color-scheme:light dark;--bg:#fcfcfb;--panel:#ffffff;--text:#1a1a19;--muted:#5f5e58;--line:#d9d8d2;--head:#eeede8;--accent:#2a78d6;
--s1:#2a78d6;--s2:#eb6834;--s3:#1baf7a;--s4:#eda100;--s5:#e87ba4;--s6:#008300;--s7:#4a3aa7;--s8:#e34948;
--crit:#b3261e;--high:#c2410c;--med:#8a6d00;--low:#1d6f42}
@media (prefers-color-scheme: dark){:root:not([data-theme="light"]){--bg:#1a1a19;--panel:#232321;--text:#f2f1ec;--muted:#c3c2b7;--line:#3d3d3a;--head:#2c2c2a;--accent:#3987e5;
--s1:#3987e5;--s2:#d95926;--s3:#199e70;--s4:#c98500;--s5:#d55181;--s6:#008300;--s7:#9085e9;--s8:#e66767;--crit:#ff7b72;--high:#ffa657;--med:#e3b341;--low:#7ee787}}
:root[data-theme="dark"]{--bg:#1a1a19;--panel:#232321;--text:#f2f1ec;--muted:#c3c2b7;--line:#3d3d3a;--head:#2c2c2a;--accent:#3987e5;
--s1:#3987e5;--s2:#d95926;--s3:#199e70;--s4:#c98500;--s5:#d55181;--s6:#008300;--s7:#9085e9;--s8:#e66767;--crit:#ff7b72;--high:#ffa657;--med:#e3b341;--low:#7ee787}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--text);font:14px/1.5 system-ui,-apple-system,"Segoe UI",Roboto,sans-serif}
main{max-width:1100px;margin:0 auto;padding:24px 16px}
header.report{border-bottom:2px solid var(--accent);margin-bottom:24px;padding-bottom:12px}
header.report h1{margin:0 0 4px;font-size:26px}header.report p{margin:0;color:var(--muted)}
section{margin:0 0 28px;break-inside:avoid}section h2{font-size:18px;margin:0 0 10px;padding-bottom:4px;border-bottom:1px solid var(--line)}
table{border-collapse:collapse;width:100%;font-size:13px;background:var(--panel)}
th,td{border:1px solid var(--line);padding:6px 8px;text-align:left;vertical-align:top;word-break:break-word}
th{background:var(--head);position:sticky;top:0}tbody tr:nth-child(even){background:color-mix(in srgb,var(--head) 40%,var(--panel))}
.kv th{width:28%}.notes{color:var(--muted);font-size:12px;margin:6px 0 0}
pre{background:var(--panel);border:1px solid var(--line);padding:10px;overflow:auto;white-space:pre-wrap;font-size:12.5px}
.badge{display:inline-block;padding:1px 7px;border-radius:9px;font-size:11.5px;font-weight:600;border:1px solid currentColor}
.sev-CRITICAL,.v-VULNERABLE{color:var(--crit)}.sev-HIGH{color:var(--high)}.sev-MEDIUM{color:var(--med)}.sev-LOW,.sev-INFO,.sev-NONE,.v-RESISTED,.v-BLOCKED{color:var(--low)}
.chart{display:flex;flex-wrap:wrap;gap:20px;align-items:flex-start}.chart svg{max-width:100%;height:auto}.chart table{width:auto;min-width:220px}
svg text{fill:var(--text);font-size:11px}svg .axis{stroke:var(--line)}svg .val{fill:var(--muted)}
.s1{fill:var(--s1)}.s2{fill:var(--s2)}.s3{fill:var(--s3)}.s4{fill:var(--s4)}.s5{fill:var(--s5)}.s6{fill:var(--s6)}.s7{fill:var(--s7)}.s8{fill:var(--s8)}
.swatch{display:inline-block;width:10px;height:10px;border-radius:2px;margin-right:6px;vertical-align:middle}
footer{color:var(--muted);font-size:12px;border-top:1px solid var(--line);margin-top:24px;padding-top:8px}
@media print{body{background:#fff;color:#000}main{max-width:none;padding:0}th{position:static}section{page-break-inside:avoid}}
"""


def _svg_bar(section: ChartSection) -> str:
    labels, values = section.labels, [float(v) for v in section.values]
    if not labels:
        return "<p class='notes'>No data.</p>"
    row_h, left, width = 22, 140, 520
    top, height = 8, row_h * len(labels) + 16
    vmax = max(values) if max(values) > 0 else 1.0
    parts = [f"<svg viewBox='0 0 {width + left + 60} {height}' width='{width + left + 60}' height='{height}' role='img' aria-label='{escape(section.title)}'>"]
    parts.append(f"<line class='axis' x1='{left}' y1='{top}' x2='{left}' y2='{height - 8}' stroke-width='1'/>")
    for i, (label, value) in enumerate(zip(labels, values, strict=False)):
        y = top + i * row_h
        w = max(2.0, (value / vmax) * width) if value > 0 else 0.0
        parts.append(f"<text x='{left - 8}' y='{y + 15}' text-anchor='end'>{escape(_truncate(str(label), 22))}</text>")
        parts.append(f"<title>{escape(str(label))}: {_cell(value)}</title>")
        parts.append(f"<rect class='s1' x='{left + 1}' y='{y + 3}' width='{w:.1f}' height='{row_h - 6}' rx='3'><title>{escape(str(label))}: {_cell(value)}</title></rect>")
        parts.append(f"<text class='val' x='{left + w + 6:.1f}' y='{y + 15}'>{_cell(value)}</text>")
    parts.append("</svg>")
    return "".join(parts)


def _svg_pie(section: ChartSection) -> str:
    labels, values = section.labels, [max(0.0, float(v)) for v in section.values]
    total = sum(values)
    if not labels or total <= 0:
        return "<p class='notes'>No data.</p>"
    cx = cy = 90
    r = 80
    parts = [f"<svg viewBox='0 0 180 180' width='180' height='180' role='img' aria-label='{escape(section.title)}'>"]
    angle = -math.pi / 2
    for i, (label, value) in enumerate(zip(labels, values, strict=False)):
        if value <= 0:
            continue
        frac = value / total
        cls = f"s{(i % 8) + 1}"
        if frac >= 0.9999:
            parts.append(f"<circle class='{cls}' cx='{cx}' cy='{cy}' r='{r}'><title>{escape(str(label))}: {_cell(value)}</title></circle>")
            break
        end = angle + frac * 2 * math.pi
        x1, y1 = cx + r * math.cos(angle), cy + r * math.sin(angle)
        x2, y2 = cx + r * math.cos(end), cy + r * math.sin(end)
        large = 1 if frac > 0.5 else 0
        parts.append(f"<path class='{cls}' d='M{cx},{cy} L{x1:.2f},{y1:.2f} A{r},{r} 0 {large},1 {x2:.2f},{y2:.2f} Z' stroke='var(--bg)' stroke-width='2'><title>{escape(str(label))}: {_cell(value)} ({frac * 100:.1f}%)</title></path>")
        angle = end
    parts.append("</svg>")
    return "".join(parts)


def _html_section(s: Section) -> str:
    body: str
    if isinstance(s, TextSection):
        body = _md_to_html(s.text) if s.markdown else f"<pre>{escape(s.text)}</pre>"
    elif isinstance(s, KeyValueSection):
        rows = "".join(f"<tr><th>{escape(str(k))}</th><td>{_html_value(v)}</td></tr>" for k, v in s.rows)
        body = f"<table class='kv'><tbody>{rows}</tbody></table>"
    elif isinstance(s, ChartSection):
        columns, rows = s.as_table()
        legend = "".join(f"<tr><td><span class='swatch s{(i % 8) + 1}'></span>{escape(_cell(r[0]))}</td><td>{escape(_cell(r[1]))}</td></tr>" for i, r in enumerate(rows)) if s.chart == "pie" else "".join(f"<tr><td>{escape(_cell(r[0]))}</td><td>{escape(_cell(r[1]))}</td></tr>" for r in rows)
        svg = _svg_pie(s) if s.chart == "pie" else _svg_bar(s)
        body = f"<div class='chart'>{svg}<table><thead><tr><th>{escape(columns[0])}</th><th>{escape(columns[1])}</th></tr></thead><tbody>{legend}</tbody></table></div>"
    else:
        columns, rows = s.as_table()
        head = "".join(f"<th>{escape(str(c))}</th>" for c in columns)
        body_rows = "".join("<tr>" + "".join(f"<td>{_html_value(c)}</td>" for c in r) + "</tr>" for r in rows) or f"<tr><td colspan='{max(1, len(columns))}' class='notes'>(no rows)</td></tr>"
        body = f"<table><thead><tr>{head}</tr></thead><tbody>{body_rows}</tbody></table>"
        if isinstance(s, TableSection) and s.notes:
            body += f"<p class='notes'>{escape(s.notes)}</p>"
    return f"<section id='{_slug(s.title)}'><h2>{escape(s.title)}</h2>{body}</section>"


def _html_value(value: Any) -> str:
    text = _cell(value)
    upper = text.upper()
    if upper in ("CRITICAL", "HIGH", "MEDIUM", "LOW", "INFO", "NONE"):
        return f"<span class='badge sev-{upper}'>{escape(text)}</span>"
    if upper in ("VULNERABLE", "RESISTED", "BLOCKED"):
        return f"<span class='badge v-{upper}'>{escape(text)}</span>"
    return escape(text)


def _md_to_html(text: str) -> str:
    """Tiny markdown subset: headings, bullet lists, code spans, bold, paragraphs."""
    html: list[str] = []
    in_list = False
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            if in_list:
                html.append("</ul>")
                in_list = False
            continue
        if line.startswith(("- ", "* ")):
            if not in_list:
                html.append("<ul>")
                in_list = True
            html.append(f"<li>{_inline_md(line[2:])}</li>")
            continue
        if in_list:
            html.append("</ul>")
            in_list = False
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level = min(6, len(m.group(1)) + 2)
            html.append(f"<h{level}>{_inline_md(m.group(2))}</h{level}>")
        else:
            html.append(f"<p>{_inline_md(line)}</p>")
    if in_list:
        html.append("</ul>")
    return "".join(html)


def _inline_md(text: str) -> str:
    out = escape(text)
    out = re.sub(r"`([^`]+)`", r"<code>\1</code>", out)
    out = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", out)
    return out


def render_html(report: Report) -> bytes:
    sections = "".join(_html_section(s) for s in report.sections)
    toc = "".join(f"<li><a href='#{_slug(s.title)}'>{escape(s.title)}</a></li>" for s in report.sections)
    doc = (
        "<!doctype html><html lang='en'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>"
        f"<title>{escape(report.title)}</title><style>{HTML_CSS}</style></head><body><main>"
        f"<header class='report'><h1>{escape(report.title)}</h1><p>{escape(report.subtitle)}</p>"
        f"<p>Generated {escape(report.generated_at.isoformat(timespec='seconds'))} by {escape(report.generated_by)} (AIRT {escape(__version__)})</p></header>"
        f"<nav><ul>{toc}</ul></nav>{sections}"
        f"<footer>AIRT {escape(__version__)} report, {escape(report.generated_at.isoformat(timespec='seconds'))}</footer></main></body></html>"
    )
    return doc.encode("utf-8")


# --- xml ----------------------------------------------------------------------------
def _xml_safe(text: Any) -> str:
    s = _cell(text)
    return "".join(ch for ch in s if ch in "\t\n\r" or 0x20 <= ord(ch) <= 0xD7FF or 0xE000 <= ord(ch) <= 0xFFFD)


def render_xml(report: Report) -> bytes:
    root = ET.Element("report", {"title": _xml_safe(report.title), "generated_at": report.generated_at.isoformat(), "generated_by": _xml_safe(report.generated_by), "version": __version__})
    ET.SubElement(root, "subtitle").text = _xml_safe(report.subtitle)
    meta = ET.SubElement(root, "meta")
    for k, v in report.meta.items():
        ET.SubElement(meta, "item", {"key": _xml_safe(k)}).text = _xml_safe(v)
    sections = ET.SubElement(root, "sections")
    for s in report.sections:
        el = ET.SubElement(sections, "section", {"kind": s.kind, "title": _xml_safe(s.title)})
        if isinstance(s, TextSection):
            el.text = _xml_safe(s.text)
            el.set("markdown", "true" if s.markdown else "false")
        elif isinstance(s, KeyValueSection):
            for k, v in s.rows:
                ET.SubElement(el, "item", {"key": _xml_safe(k)}).text = _xml_safe(v)
        elif isinstance(s, FindingsSection):
            for f in s.findings:
                fe = ET.SubElement(el, "finding")
                for k, v in f.items():
                    ET.SubElement(fe, "field", {"name": _xml_safe(k)}).text = _xml_safe(v)
        elif isinstance(s, ChartSection):
            el.set("chart", s.chart)
            for label, value in zip(s.labels, s.values, strict=False):
                ET.SubElement(el, "point", {"label": _xml_safe(label)}).text = _cell(value)
        else:
            columns, rows = s.as_table()
            cols = ET.SubElement(el, "columns")
            for c in columns:
                ET.SubElement(cols, "column").text = _xml_safe(c)
            rows_el = ET.SubElement(el, "rows")
            for r in rows:
                re_ = ET.SubElement(rows_el, "row")
                for c in r:
                    ET.SubElement(re_, "cell").text = _xml_safe(c)
            if isinstance(s, TableSection) and s.notes:
                ET.SubElement(el, "notes").text = _xml_safe(s.notes)
    ET.indent(root)
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


# --- sarif --------------------------------------------------------------------------
def render_sarif(report: Report) -> bytes:
    findings = report.findings()
    rules: dict[str, dict[str, Any]] = {}
    results: list[dict[str, Any]] = []
    for f in findings:
        category = str(f.get("category") or "unknown")
        severity = str(f.get("severity") or "MEDIUM").upper()
        level = SEVERITY_TO_SARIF.get(severity, "warning")
        if category not in rules:
            rules[category] = {
                "id": category,
                "name": category.replace("_", " ").title().replace(" ", ""),
                "shortDescription": {"text": f"AIRT finding category: {category}"},
                "fullDescription": {"text": f"Findings raised by AIRT analyzers or red-team probes in the '{category}' category."},
                "defaultConfiguration": {"level": level},
                "properties": {"tags": ["security", "llm", category]},
            }
        message = str(f.get("message") or f.get("title") or category)
        if f.get("description"):
            message = f"{message}: {f['description']}"
        props = {k: _plain(v) for k, v in f.items() if k not in ("title", "message", "description") and v not in (None, "")}
        props["severity"] = severity
        result: dict[str, Any] = {
            "ruleId": category,
            "ruleIndex": list(rules).index(category),
            "level": level,
            "message": {"text": message},
            "properties": props,
        }
        if f.get("evidence"):
            result["message"]["markdown"] = f"{message}\n\nEvidence: `{_truncate(str(f['evidence']), 300)}`"
        location = str(f.get("location") or f.get("probe_id") or f.get("ticket_id") or "")
        if location:
            result["locations"] = [{"logicalLocations": [{"name": location, "kind": "member"}]}]
        fingerprint_parts = [category, str(f.get("ticket_id") or ""), str(f.get("probe_id") or ""), str(f.get("title") or "")]
        result["partialFingerprints"] = {"airt/v1": "|".join(fingerprint_parts)}
        results.append(result)
    doc = {
        "$schema": "https://json.schemastore.org/sarif-2.1.0.json",
        "version": "2.1.0",
        "runs": [
            {
                "tool": {"driver": {"name": "AIRT", "fullName": "AIRT AI Red Team Gateway", "version": __version__, "informationUri": "https://github.com/airt", "rules": list(rules.values())}},
                "automationDetails": {"id": f"airt/{report.meta.get('kind', 'report')}", "description": {"text": report.title}},
                "invocations": [{"executionSuccessful": True, "endTimeUtc": report.generated_at.astimezone(UTC).strftime("%Y-%m-%dT%H:%M:%SZ")}],
                "results": results,
                "properties": {"title": report.title, "subtitle": report.subtitle, "generated_by": report.generated_by, "meta": _report_json(report)["meta"]},
            }
        ],
    }
    return json.dumps(doc, indent=2, default=str).encode("utf-8")


# --- junit --------------------------------------------------------------------------
def _pick_column(columns: list[str], names: list[str]) -> int | None:
    lowered = [c.lower() for c in columns]
    for n in names:
        if n in lowered:
            return lowered.index(n)
    return None


def _junit_cases_from_table(s: TableSection) -> list[dict[str, Any]]:
    columns, rows = s.as_table()
    name_idx = _pick_column(columns, ["probe", "ticket", "#", "id", "name"])
    if name_idx is None:
        name_idx = 0
    cases: list[dict[str, Any]] = []
    for r in rows:
        values = {columns[i].lower(): _cell(r[i]) for i in range(min(len(columns), len(r)))}
        name = _cell(r[name_idx]) if r else "row"
        detail = ", ".join(f"{columns[i]}={_cell(r[i])}" for i in range(min(len(columns), len(r))))
        flagged = [v.upper() for k, v in values.items() if k in ("verdict", "risk", "risk_level", "severity", "status", "level")]
        status = "passed"
        message = ""
        if any(v in FAIL_VALUES for v in flagged):
            status = "failed"
            message = "; ".join(f"{k}={v}" for k, v in values.items() if v.upper() in FAIL_VALUES)
        elif any(v in ERROR_VALUES for v in flagged):
            status = "error"
            message = "; ".join(f"{k}={v}" for k, v in values.items() if v.upper() in ERROR_VALUES)
        elif "pending" in [v.lower() for v in flagged]:
            status = "skipped"
        cases.append({"name": name, "classname": s.title, "status": status, "message": message, "detail": detail, "time": values.get("latency ms") or values.get("latency_ms") or ""})
    return cases


def _junit_cases_from_findings(s: FindingsSection) -> list[dict[str, Any]]:
    cases: list[dict[str, Any]] = []
    for f in s.findings:
        severity = str(f.get("severity") or "").upper()
        verdict = str(f.get("verdict") or "").upper()
        failed = severity in ("CRITICAL", "HIGH") or verdict == "VULNERABLE"
        name = f"{f.get('category', 'finding')}: {f.get('title', '')}"
        if f.get("ticket_id") or f.get("probe_id"):
            name += f" [{f.get('probe_id') or f.get('ticket_id')}]"
        detail = json.dumps({k: _plain(v) for k, v in f.items()}, indent=1, default=str)
        cases.append({"name": name, "classname": s.title, "status": "failed" if failed else "passed", "message": f"{severity} {f.get('category', '')} {verdict}".strip(), "detail": detail, "time": ""})
    return cases


def render_junit(report: Report) -> bytes:
    suites: list[tuple[str, list[dict[str, Any]]]] = []
    case_tables = [s for s in report.sections if isinstance(s, TableSection) and s.role == "cases"]
    if case_tables:
        suites = [(s.title, _junit_cases_from_table(s)) for s in case_tables]
    else:
        suites = [(s.title, _junit_cases_from_findings(s)) for s in report.sections if isinstance(s, FindingsSection)]
    if not suites or not any(cases for _, cases in suites):
        suites = [("report", [{"name": report.title, "classname": "report", "status": "passed", "message": "", "detail": "no testcases derived from this report", "time": ""}])]
    root = ET.Element("testsuites", {"name": _xml_safe(report.title)})
    totals = {"tests": 0, "failures": 0, "errors": 0, "skipped": 0}
    ts = report.generated_at.astimezone(UTC).strftime("%Y-%m-%dT%H:%M:%S")
    for title, cases in suites:
        suite = ET.SubElement(root, "testsuite", {"name": _xml_safe(title), "timestamp": ts, "hostname": "airt"})
        counts = {"tests": len(cases), "failures": 0, "errors": 0, "skipped": 0}
        for c in cases:
            secs = "0"
            try:
                secs = f"{float(c['time']) / 1000.0:.3f}" if c["time"] else "0"
            except ValueError:
                secs = "0"
            tc = ET.SubElement(suite, "testcase", {"name": _xml_safe(c["name"]), "classname": _xml_safe(f"airt.{_slug(report.meta.get('kind', 'report'))}.{_slug(c['classname'])}"), "time": secs})
            if c["status"] == "failed":
                counts["failures"] += 1
                ET.SubElement(tc, "failure", {"message": _xml_safe(c["message"]), "type": "AssertionError"}).text = _xml_safe(c["detail"])
            elif c["status"] == "error":
                counts["errors"] += 1
                ET.SubElement(tc, "error", {"message": _xml_safe(c["message"]), "type": "Error"}).text = _xml_safe(c["detail"])
            elif c["status"] == "skipped":
                counts["skipped"] += 1
                ET.SubElement(tc, "skipped", {"message": "pending"})
            else:
                ET.SubElement(tc, "system-out").text = _xml_safe(c["detail"])
        for k, v in counts.items():
            suite.set(k, str(v))
            totals[k] += v
        suite.set("time", "0")
    for k, v in totals.items():
        root.set(k, str(v))
    root.set("time", "0")
    ET.indent(root)
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


# --- xlsx ---------------------------------------------------------------------------
def render_xlsx(report: Report) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2A78D6")
    used: set[str] = set()

    def sheet_name(title: str) -> str:
        base = re.sub(r"[\[\]:*?/\\]", " ", title).strip()[:28] or "Sheet"
        name = base
        n = 2
        while name.lower() in used:
            name = f"{base[:25]} {n}"
            n += 1
        used.add(name.lower())
        return name

    def write_table(ws: Any, columns: list[str], rows: list[list[Any]], start_row: int = 1) -> None:
        ws.append(columns)
        for cell in ws[start_row]:
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for r in rows:
            ws.append([_plain(c) for c in r])
        ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
        for i, col in enumerate(columns, start=1):
            longest = max([len(str(col))] + [len(_cell(r[i - 1])) for r in rows if i - 1 < len(r)]) if rows else len(str(col))
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, longest + 2))
        ws.auto_filter.ref = ws.dimensions

    cover = wb.create_sheet(sheet_name("Report"))
    write_table(cover, ["Key", "Value"], [["Title", report.title], ["Subtitle", report.subtitle], ["Generated at", report.generated_at.isoformat()], ["Generated by", report.generated_by], *[[k, _plain(v)] for k, v in report.meta.items()]])
    for s in report.sections:
        ws = wb.create_sheet(sheet_name(s.title))
        if isinstance(s, TextSection):
            write_table(ws, ["Text"], [[line] for line in s.text.splitlines()] or [[""]])
            ws.column_dimensions["A"].width = 100
            continue
        columns, rows = s.as_table()
        write_table(ws, columns or ["(empty)"], rows)
        if isinstance(s, TableSection) and s.notes:
            ws.append([])
            ws.append([s.notes])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- pdf ----------------------------------------------------------------------------
def render_pdf(report: Report) -> bytes:
    from reportlab.graphics.charts.barcharts import HorizontalBarChart
    from reportlab.graphics.charts.piecharts import Pie
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    wide = any(len(s.as_table()[0]) > 6 for s in report.sections if isinstance(s, (TableSection, FindingsSection)))
    pagesize = landscape(A4) if wide else A4
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=8.5, leading=11, alignment=TA_LEFT)
    cell_style = ParagraphStyle("cell", parent=body, fontSize=7.5, leading=9.5)
    head_style = ParagraphStyle("head", parent=cell_style, textColor=colors.white, fontName="Helvetica-Bold")
    mono = ParagraphStyle("mono", parent=body, fontName="Courier", fontSize=7.5, leading=9.5)
    title_style = styles["Title"]
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], spaceBefore=10, spaceAfter=4, textColor=colors.HexColor("#1c5cab"))
    subtitle_style = ParagraphStyle("sub", parent=body, textColor=colors.HexColor("#5f5e58"), fontSize=9.5)
    avail_width = pagesize[0] - 30 * mm

    def para(text: Any, style: ParagraphStyle = cell_style, limit: int = 1200) -> Paragraph:
        s = escape(_truncate(_cell(text), limit)).replace("\n", "<br/>")
        return Paragraph(s, style)

    def table(columns: list[str], rows: list[list[Any]]) -> Table:
        if not columns:
            return Table([[para("(empty)")]])
        weights = []
        for i, c in enumerate(columns):
            longest = max([len(str(c))] + [len(_cell(r[i])) for r in rows[:200] if i < len(r)])
            weights.append(min(60, max(6, longest)))
        total = float(sum(weights))
        widths = [avail_width * w / total for w in weights]
        data: list[list[Any]] = [[para(c, head_style) for c in columns]]
        for r in rows:
            padded = list(r) + [""] * (len(columns) - len(r))
            data.append([para(v) for v in padded[: len(columns)]])
        if not rows:
            data.append([para("(no rows)")] + [para("") for _ in columns[1:]])
        t = Table(data, colWidths=widths, repeatRows=1)
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2a78d6")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#c9c8c2")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f2ee")]),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]
        t.setStyle(TableStyle(style))
        return t

    def chart(section: ChartSection) -> Drawing:
        labels = [str(x) for x in section.labels]
        values = [float(v) for v in section.values]
        if not labels or not values or sum(values) <= 0:
            d = Drawing(300, 30)
            d.add(String(0, 10, "No data.", fontSize=9))
            return d
        if section.chart == "pie":
            d = Drawing(avail_width, 190)
            pie = Pie()
            pie.x, pie.y, pie.width, pie.height = 150, 15, 160, 160
            pie.data = values
            pie.labels = [f"{label} ({_cell(v)})" for label, v in zip(labels, values, strict=False)]
            pie.sideLabels = True
            pie.slices.strokeColor = colors.white
            pie.slices.strokeWidth = 1
            pie.slices.fontSize = 7.5
            pie.slices.fontName = "Helvetica"
            for i in range(len(values)):
                pie.slices[i].fillColor = colors.HexColor(PALETTE_LIGHT[i % len(PALETTE_LIGHT)])
            d.add(pie)
            return d
        height = max(60, 18 * len(labels) + 30)
        d = Drawing(avail_width, height)
        bc = HorizontalBarChart()
        bc.x, bc.y, bc.width, bc.height = 120, 12, avail_width - 180, height - 24
        bc.data = [values]
        bc.categoryAxis.categoryNames = [_truncate(x, 24) for x in labels]
        bc.categoryAxis.labels.fontSize = 7.5
        bc.categoryAxis.labels.fontName = "Helvetica"
        bc.categoryAxis.tickLeft = 0
        bc.categoryAxis.tickRight = 0
        bc.categoryAxis.strokeColor = colors.HexColor("#c9c8c2")
        bc.categoryAxis.labels.boxAnchor = "e"
        bc.categoryAxis.labels.dx = -4
        bc.categoryAxis.reverseDirection = True
        bc.valueAxis.valueMin = 0
        bc.valueAxis.labels.fontSize = 7
        bc.valueAxis.labels.fontName = "Helvetica"
        bc.valueAxis.visibleGrid = True
        bc.valueAxis.gridStrokeColor = colors.HexColor("#e6e5e0")
        bc.valueAxis.strokeColor = colors.HexColor("#c9c8c2")
        bc.bars[0].fillColor = colors.HexColor(PALETTE_LIGHT[0])
        bc.bars[0].strokeColor = None
        bc.barLabels.fontSize = 7
        bc.barLabels.fontName = "Helvetica"
        bc.barLabelFormat = lambda v: _cell(v)
        bc.barLabels.boxAnchor = "w"
        bc.barLabels.dx = 3
        d.add(bc)
        return d

    story: list[Any] = [Paragraph(escape(report.title), title_style)]
    if report.subtitle:
        story.append(Paragraph(escape(report.subtitle), subtitle_style))
    story.append(Paragraph(escape(f"Generated {report.generated_at.isoformat(timespec='seconds')} by {report.generated_by}, AIRT {__version__}"), subtitle_style))
    story.append(Spacer(1, 8))
    for s in report.sections:
        block: list[Any] = [Paragraph(escape(s.title), h2)]
        if isinstance(s, TextSection):
            for chunk in (s.text or "(empty)").split("\n\n"):
                block.append(para(chunk, mono if not s.markdown else body, 4000))
        elif isinstance(s, KeyValueSection):
            block.append(table(["Key", "Value"], [[k, v] for k, v in s.rows]))
        elif isinstance(s, ChartSection):
            block.append(chart(s))
            block.append(table(*s.as_table()))
        else:
            columns, rows = s.as_table()
            block.append(table(columns, rows))
            if isinstance(s, TableSection) and s.notes:
                block.append(para(s.notes, subtitle_style))
        small = len(getattr(s, "rows", [])) < 12
        if small:
            story.append(KeepTogether(block))
        else:
            story.extend(block)
        story.append(Spacer(1, 6))

    header_text = f"{report.title} - {report.generated_at.strftime('%Y-%m-%d %H:%M UTC')}"

    def on_page(canvas: Any, doc: Any) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#5f5e58"))
        canvas.drawString(15 * mm, pagesize[1] - 10 * mm, _truncate(header_text, 120))
        canvas.drawRightString(pagesize[0] - 15 * mm, pagesize[1] - 10 * mm, f"AIRT {__version__}")
        canvas.setStrokeColor(colors.HexColor("#c9c8c2"))
        canvas.line(15 * mm, pagesize[1] - 11.5 * mm, pagesize[0] - 15 * mm, pagesize[1] - 11.5 * mm)
        canvas.drawCentredString(pagesize[0] / 2, 8 * mm, f"Page {doc.page}")
        canvas.drawString(15 * mm, 8 * mm, "Confidential - generated by AIRT")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=pagesize, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=16 * mm, bottomMargin=14 * mm, title=report.title, author=report.generated_by, subject=report.subtitle)
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()


# --- dispatch -----------------------------------------------------------------------
RENDERERS: dict[str, Callable[[Report], bytes]] = {
    "json": render_json,
    "yaml": render_yaml,
    "csv": render_csv,
    "tsv": render_tsv,
    "md": render_md,
    "html": render_html,
    "pdf": render_pdf,
    "xlsx": render_xlsx,
    "txt": render_txt,
    "xml": render_xml,
    "sarif": render_sarif,
    "junit": render_junit,
}


def render(report: Report, fmt: str) -> tuple[bytes, str, str]:
    key = normalize_format(fmt)
    payload = RENDERERS[key](report)
    info = FORMATS[key]
    return payload, info["media_type"], info["extension"]
